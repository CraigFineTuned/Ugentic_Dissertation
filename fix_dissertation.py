import os
from datetime import datetime
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Define paths
base_dir = r"C:\Users\craig\Desktop\MainProjects\Ugentic_Dissertation\DISSERTATION_ACADEMIC\Completed"
input_file = os.path.join(base_dir, "Final_Dissertation_YOLO.docx")
output_file = os.path.join(base_dir, "SUBMISSION_READY_Dissertation.docx")
checklist_file = os.path.join(base_dir, "Comprehensive Checklist.md")

def fix_document(doc_path, save_path):
    if not os.path.exists(doc_path):
        print(f"Error: File not found at {doc_path}")
        return

    print(f"Reading document: {doc_path}")
    doc = Document(doc_path)

    # --- Apply Formatting Standards (The "Polish") ---
    # We will iterate through paragraphs to apply direct formatting.
    # Note: Modifying styles is cleaner, but direct formatting ensures it 'sticks' over any local overrides.
    
    # 1. Update Normal Style first as a baseline
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 1.5
    
    # 2. Iterate through all paragraphs to enforce Justification and Spacing
    count = 0
    for paragraph in doc.paragraphs:
        count += 1
        # SKIP Heading styles for Justification (Headings are usually Left or Centered)
        if paragraph.style.name.startswith('Heading'):
            continue
        
        # SKIP Centered paragraphs (like Figure placeholders or Titles)
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            continue
            
        # Enforce Justified Alignment for body text
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        # Enforce 1.5 Line Spacing (redundant if style is set, but safe)
        paragraph.paragraph_format.line_spacing = 1.5

    print(f"Processed {count} paragraphs.")

    # --- Save the Fixed Document ---
    doc.save(save_path)
    print(f"✅ Document fixed and saved to: {save_path}")

    # --- Update Checklist ---
    update_checklist(checklist_file)

def update_checklist(md_path):
    if not os.path.exists(md_path):
        print("Checklist file not found, skipping update.")
        return

    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    with open(md_path, "a") as f:
        f.write("\n\n## AUTOMATED VERIFICATION LOG (Phase 2 Fix)\n")
        f.write(f"* **Date:** {timestamp}\n")
        f.write("* **Status:** Document processed with 'Final_Dissertation_YOLO.docx' as source.\n")
        f.write("* **Formatting Applied:**\n")
        f.write("    * Font: Times New Roman, 12pt\n")
        f.write("    * Alignment: Justified (Body text), Left/Center preserved for Headings\n")
        f.write("    * Spacing: 1.5 Line Spacing global\n")
        f.write("* **Next Steps:** User must manually insert images and ethical clearance letter.\n")

    print(f"✅ Checklist updated at: {md_path}")

if __name__ == "__main__":
    print("Starting Dissertation Fixer...")
    fix_document(input_file, output_file)
